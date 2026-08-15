#!/usr/bin/env python3
"""
Converts Data/YMR.xlsx into assets/data.js so the dashboard (dashboard.html)
can load it with a plain <script src="assets/data.js"> tag — no server and no
browser file:// fetch restrictions required.

Run this any time the source Excel file changes, then reload dashboard.html
in Chrome (or just double-click Update_Dashboard.bat, which does both).
"""
import datetime
import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("ERROR: the 'openpyxl' package is required.")
    print("Install it once with:  pip install openpyxl")
    sys.exit(1)

ROOT = Path(__file__).resolve().parent.parent
SOURCE_FILE = ROOT / "Data" / "YMR.xlsx"
SHEET_NAME = "YMR Factors"
OUTPUT_FILE = ROOT / "assets" / "data.js"

# Maps the 5 attributes the dashboard tracks (per instruction.md checkboxes)
# to the "Attribute" label used in the source file's M:Q table, and to the
# camelCase field name used for the matching data column (D:K).
ATTRIBUTE_MAP = {
    "Shortfall %":     {"key": "shortfall",   "checkboxLabel": "Shortfall"},
    "C-Grade %":       {"key": "cGrade",      "checkboxLabel": "C Grade"},
    "Yarn Waste %":    {"key": "yarnWaste",   "checkboxLabel": "Yarn Waste"},
    "UC Small Lots %": {"key": "ucSmallLots", "checkboxLabel": "UC Small Lot"},
    "Gain/Loss %":     {"key": "gainLoss",    "checkboxLabel": "Gain/Loss"},
}

# Source column header (row A:K) -> camelCase field name used everywhere else.
FIELD_MAP = {
    "Month": "month",
    "Plant": "plant",
    "Unit": "unit",
    "Consumption": "consumption",
    "Yield": "yield",
    "Total Waste": "totalWaste",
    "Shortfall": "shortfall",
    "C-Grade": "cGrade",
    "Yarn Waste": "yarnWaste",
    "UC Small Lots": "ucSmallLots",
    "Gain/Loss": "gainLoss",
}


def find_header_row(ws, marker_col_letter, marker_value, max_scan=20):
    for row in ws.iter_rows(min_row=1, max_row=max_scan):
        for cell in row:
            if cell.column_letter == marker_col_letter and cell.value == marker_value:
                return cell.row
    raise ValueError(f"Could not find header cell {marker_col_letter}='{marker_value}'")


def num(v):
    if v is None:
        return 0
    return float(v)


def main():
    if not SOURCE_FILE.exists():
        print(f"ERROR: source file not found: {SOURCE_FILE}")
        sys.exit(1)

    wb = openpyxl.load_workbook(SOURCE_FILE, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        print(f"ERROR: sheet '{SHEET_NAME}' not found. Sheets present: {wb.sheetnames}")
        sys.exit(1)
    ws = wb[SHEET_NAME]

    # ---- Data table (A:K) ----
    data_header_row = find_header_row(ws, "A", "Month")
    header_cells = ws[data_header_row]
    col_to_field = {}
    for cell in header_cells:
        field = FIELD_MAP.get(cell.value)
        if field:
            col_to_field[cell.column] = field

    rows = []
    r = data_header_row + 1
    while True:
        month_val = ws.cell(row=r, column=1).value
        if month_val is None:
            # allow a single blank row inside the table, but stop after two in a row
            next_val = ws.cell(row=r + 1, column=1).value
            if next_val is None:
                break
            r += 1
            continue
        record = {}
        for col, field in col_to_field.items():
            val = ws.cell(row=r, column=col).value
            record[field] = val if field in ("month", "plant", "unit") else num(val)
        rows.append(record)
        r += 1

    # ---- Attribute / target table (M:Q) ----
    attr_header_row = find_header_row(ws, "M", "Attribute")
    attributes = []
    r = attr_header_row + 1
    found_keys = set()
    while len(found_keys) < len(ATTRIBUTE_MAP):
        label = ws.cell(row=r, column=13).value  # M
        if label is None:
            r += 1
            if r > attr_header_row + 40:
                break
            continue
        if label in ATTRIBUTE_MAP:
            meta = ATTRIBUTE_MAP[label]
            formula = ws.cell(row=r, column=14).value  # N
            target_p05 = ws.cell(row=r, column=15).value  # O
            target_p06 = ws.cell(row=r, column=16).value  # P
            direction = ws.cell(row=r, column=17).value  # Q
            attributes.append({
                "key": meta["key"],
                "label": meta["checkboxLabel"],
                "sourceLabel": label,
                "formula": formula,
                "direction": direction,
                "targets": {
                    "Plant-05": num(target_p05),
                    "Plant-06": num(target_p06),
                },
            })
            found_keys.add(label)
        r += 1

    missing = set(ATTRIBUTE_MAP) - found_keys
    if missing:
        print(f"WARNING: could not find attribute definitions for: {sorted(missing)}")

    # Keep attributes in the fixed spec order regardless of source row order.
    order = ["shortfall", "cGrade", "yarnWaste", "ucSmallLots", "gainLoss"]
    attributes.sort(key=lambda a: order.index(a["key"]))

    payload = {
        "generatedAt": datetime.datetime.now().isoformat(timespec="seconds"),
        "sourceFile": "Data/YMR.xlsx",
        "rows": rows,
        "attributes": attributes,
    }

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    js = (
        "// AUTO-GENERATED by scripts/refresh_data.py — do not edit by hand.\n"
        "// Re-run the script (or double-click Update_Dashboard.bat) after changing Data/YMR.xlsx.\n"
        "const YMR_DATA = " + json.dumps(payload, indent=2) + ";\n"
    )
    OUTPUT_FILE.write_text(js, encoding="utf-8")
    print(f"OK: wrote {len(rows)} data rows and {len(attributes)} attributes to {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
